"""Scrape all non-WordPress kendo sources: kendo-guide.com, official docs,
PDFs, Vietnamese sites, and other single-page sources.

These don't fit the base_scraper WordPress pattern, so they use direct
requests + lxml extraction.

Usage:
    python scripts/scraping/scrape_all_new.py
    python scripts/scraping/scrape_all_new.py --dry-run
    python scripts/scraping/scrape_all_new.py --only pdfs
    python scripts/scraping/scrape_all_new.py --only kendoguide
    python scripts/scraping/scrape_all_new.py --only official
    python scripts/scraping/scrape_all_new.py --only vietnamese
"""

from __future__ import annotations

import argparse
import re
import sys
import time
from pathlib import Path
from urllib.parse import urljoin

_project_root = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(_project_root / "src"))

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

import requests
import yaml
from lxml import html

from kendocenter.config import settings

THEORY_DIR = settings.theory_path
HEADERS = {
    "User-Agent": "Mozilla/5.0 (compatible; YSKKenjin/1.0; kendo knowledge base)"
}


def slugify(text: str) -> str:
    text = text.strip().lower()
    text = re.sub(r"[^\w\s-]", "", text)
    text = re.sub(r"[\s_]+", "-", text)
    text = re.sub(r"-+", "-", text)
    return text[:80].strip("-")


def save_docx(title: str, paragraphs: list[str], output_path: Path,
              url: str = "", date: str = "", source: str = "") -> None:
    from docx import Document
    doc = Document()
    doc.add_heading(title, level=1)
    if date:
        doc.add_paragraph(f"Date: {date}")
    if source:
        doc.add_paragraph(f"Source: {source}")
    if url:
        doc.add_paragraph(f"URL: {url}")
    doc.add_paragraph("")
    for p in paragraphs:
        doc.add_paragraph(p)
    doc.save(str(output_path))


def fetch_page(url: str) -> html.HtmlElement | None:
    try:
        resp = requests.get(url, headers=HEADERS, timeout=20)
        resp.raise_for_status()
        return html.fromstring(resp.text)
    except Exception as e:
        print(f"    ERROR fetching {url}: {e}")
        return None


def extract_text_elements(tree, content_xpath: str,
                          elements_xpath: str = ".//p | .//h2 | .//h3 | .//h4 | .//li | .//div[not(div)]",
                          stop_texts: list[str] | None = None) -> list[str]:
    """Extract paragraphs from a content div."""
    stop_texts = stop_texts or []
    content_el = tree.xpath(content_xpath)
    if not content_el:
        return []
    paragraphs = []
    elements = content_el[0].xpath(elements_xpath)
    for el in elements:
        text = el.text_content().strip()
        if not text or len(text) < 5:
            continue
        if any(text.startswith(s) for s in stop_texts):
            break
        paragraphs.append(text)
    return paragraphs


def write_metadata_yaml(folder: Path, category: str, description: str,
                        language: str, files_meta: dict) -> None:
    meta_path = folder / "metadata.yaml"
    meta = {
        "category": category,
        "description": description,
        "doc_type": "article",
        "default_language": language,
        "files": files_meta,
    }
    with open(meta_path, "w", encoding="utf-8") as f:
        yaml.dump(meta, f, allow_unicode=True, default_flow_style=False, sort_keys=False)
    print(f"  Written: {meta_path}")


# ============================================================
# KENDO-GUIDE.COM (custom HTML, not WordPress)
# ============================================================

KENDOGUIDE_PAGES = [
    ("what_is_kendo.html", "What is Kendo"),
    ("kendo_history.html", "Kendo History"),
    ("kendo-philosophy.html", "Kendo Philosophy"),
    ("kendo_basics.html", "Kendo Basics"),
    ("chudan_no_kamae.html", "Chudan no Kamae"),
    ("kendo-footwork-ashi-sabaki.html", "Kendo Footwork (Ashi Sabaki)"),
    ("men_strike.html", "Men Strike"),
    ("kote_strike.html", "Kote Strike"),
    ("do_strike.html", "Do Strike"),
    ("tsuki.html", "Tsuki"),
    ("kendo_techniques.html", "Kendo Techniques Overview"),
    ("kendo_kata.html", "Kendo Kata"),
    ("terminology_seme_san_sappo.html", "Seme and San Sappo"),
    ("different-distances-in-kendo.html", "Different Distances (Maai)"),
    ("difference-and-purpose-betwen-each-type-of-keiko.html", "Types of Keiko"),
    ("requirements-to-pass-1dan-2dan-3dan.html", "Requirements to Pass Dan Exams"),
    ("kendo-suburi.html", "Kendo Suburi"),
    ("tenouchi.html", "Tenouchi"),
    ("kendo_etiquette.html", "Kendo Etiquette (Reigi)"),
    ("shinai.html", "Shinai"),
    ("kendo-bogu.html", "Kendo Bogu"),
    ("shinai-maintenance.html", "Shinai Maintenance"),
    ("kiai.html", "Kiai"),
    ("zanshin.html", "Zanshin"),
    ("kendo-sparring-jigeiko.html", "Jigeiko (Free Sparring)"),
    ("kakari-geiko.html", "Kakarigeiko"),
    ("nito.html", "Nito (Two-Sword Style)"),
    ("jodan.html", "Jodan no Kamae"),
    ("kendo-warm-up.html", "Kendo Warm-Up"),
    ("kendo_rank.html", "Kendo Ranking System"),
    ("kendo_tournament.html", "Kendo Tournament"),
    ("musashi_miyamoto.html", "Miyamoto Musashi"),
    ("wkc.html", "World Kendo Championships"),
    ("kendo-dojo.html", "Kendo Dojo"),
    ("bokken.html", "Bokken"),
    ("hakama.html", "Hakama"),
    ("keikogi.html", "Keikogi"),
]


def scrape_kendoguide(dry_run: bool = False) -> int:
    print("\n=== Scraping kendo-guide.com ===")
    folder = THEORY_DIR / "blogs" / "kendoguide"
    folder.mkdir(parents=True, exist_ok=True)
    files_meta = {}
    count = 0

    for page, title in KENDOGUIDE_PAGES:
        slug = slugify(title)
        filename = f"{slug}.docx"
        output_path = folder / filename
        url = f"https://www.kendo-guide.com/{page}"

        if output_path.exists():
            print(f"  SKIP (exists): {filename}")
            continue
        if dry_run:
            print(f"  WOULD SCRAPE: {title}")
            continue

        print(f"  Scraping: {title}...")
        tree = fetch_page(url)
        if tree is None:
            continue

        # kendo-guide.com uses various content containers
        paragraphs = []
        for xpath in [
            '//div[@id="content"]',
            '//div[contains(@class, "entry-content")]',
            '//div[contains(@class, "content")]',
            '//article',
            '//main',
        ]:
            paragraphs = extract_text_elements(tree, xpath)
            if paragraphs:
                break

        if not paragraphs:
            # Fallback: grab all p tags from body
            all_p = tree.xpath("//p")
            paragraphs = [p.text_content().strip() for p in all_p
                          if p.text_content().strip() and len(p.text_content().strip()) > 20]

        if not paragraphs:
            print(f"    WARNING: No content for {title}")
            continue

        save_docx(title, paragraphs, output_path, url=url,
                  source="kendo-guide.com by Masahiro Imafuji")
        print(f"    Saved: {filename} ({len(paragraphs)} paragraphs)")
        files_meta[filename] = {
            "title": title,
            "url": url,
            "tags": ["blog", "english", "kendoguide", "instruction"],
        }
        count += 1
        time.sleep(2.0)

    if not dry_run and files_meta:
        write_metadata_yaml(folder, "kendoguide",
                            "Kendo instruction by Masahiro Imafuji (7-dan)",
                            "en", files_meta)
    return count


# ============================================================
# OFFICIAL DOCUMENTS (single-page sources)
# ============================================================

OFFICIAL_PAGES = [
    {
        "url": "https://www.kendo.or.jp/en/knowledge/kendo-concept/",
        "title": "The Concept of Kendo (AJKF)",
        "xpath": '//div[contains(@class, "entry-content")] | //div[contains(@class, "content")] | //article | //main',
        "tags": ["official", "ajkf", "philosophy", "concept"],
    },
    {
        "url": "https://www.kendo.or.jp/en/knowledge/kendo-history/",
        "title": "History of Kendo (AJKF)",
        "xpath": '//div[contains(@class, "entry-content")] | //div[contains(@class, "content")] | //article | //main',
        "tags": ["official", "ajkf", "history"],
    },
    {
        "url": "https://www.kendo.or.jp/en/knowledge/kendo-origin/",
        "title": "AJKF Perspective on Kendo Origin",
        "xpath": '//div[contains(@class, "entry-content")] | //div[contains(@class, "content")] | //article | //main',
        "tags": ["official", "ajkf", "history", "origin"],
    },
    {
        "url": "https://www.britishkendoassociation.com/grading-questions-kendo/",
        "title": "BKA Kendo Grading Questions",
        "xpath": '//div[contains(@class, "entry-content")] | //div[contains(@class, "content")] | //article | //main',
        "tags": ["official", "bka", "grading", "questions"],
    },
    {
        "url": "https://www.britishkendoassociation.com/grading-requirements/",
        "title": "BKA Grading Requirements",
        "xpath": '//div[contains(@class, "entry-content")] | //div[contains(@class, "content")] | //article | //main',
        "tags": ["official", "bka", "grading"],
    },
    {
        "url": "https://www.britishkendoassociation.com/a-guide-for-dojo-leaders-and-examiners/",
        "title": "BKA Guide for Dojo Leaders and Examiners",
        "xpath": '//div[contains(@class, "entry-content")] | //div[contains(@class, "content")] | //article | //main',
        "tags": ["official", "bka", "grading", "instruction"],
    },
    {
        "url": "https://www.auskf.org/info/kendo-promotional-exam-study-guide",
        "title": "AUSKF Kendo Promotional Exam Study Guide",
        "xpath": '//div[contains(@class, "entry-content")] | //div[contains(@class, "content")] | //article | //main',
        "tags": ["official", "auskf", "grading", "study-guide"],
    },
    {
        "url": "https://kendoklubben.se/the-mindset-of-kendo-instruction/",
        "title": "The Mindset of Kendo Instruction (AJKF)",
        "xpath": '//div[contains(@class, "entry-content")] | //div[contains(@class, "content")] | //article',
        "tags": ["official", "ajkf", "philosophy", "instruction"],
    },
    {
        "url": "https://en.motenas-japan.jp/kendo-history/",
        "title": "History of Kendo - From Heian to Modern Times",
        "xpath": '//div[contains(@class, "entry-content")] | //div[contains(@class, "content")] | //article | //main',
        "tags": ["history", "comprehensive"],
    },
]


def scrape_official(dry_run: bool = False) -> int:
    print("\n=== Scraping official documents ===")
    folder = THEORY_DIR / "official"
    folder.mkdir(parents=True, exist_ok=True)
    files_meta = {}
    count = 0

    for page in OFFICIAL_PAGES:
        slug = slugify(page["title"])
        filename = f"{slug}.docx"
        output_path = folder / filename

        if output_path.exists():
            print(f"  SKIP (exists): {filename}")
            continue
        if dry_run:
            print(f"  WOULD SCRAPE: {page['title']}")
            continue

        print(f"  Scraping: {page['title']}...")
        tree = fetch_page(page["url"])
        if tree is None:
            continue

        # Try each xpath option
        paragraphs = []
        for xpath in page["xpath"].split(" | "):
            paragraphs = extract_text_elements(tree, xpath)
            if paragraphs:
                break

        if not paragraphs:
            all_p = tree.xpath("//p")
            paragraphs = [p.text_content().strip() for p in all_p
                          if p.text_content().strip() and len(p.text_content().strip()) > 20]

        if not paragraphs:
            print(f"    WARNING: No content for {page['title']}")
            continue

        save_docx(page["title"], paragraphs, output_path,
                  url=page["url"])
        print(f"    Saved: {filename} ({len(paragraphs)} paragraphs)")
        files_meta[filename] = {
            "title": page["title"],
            "url": page["url"],
            "tags": page["tags"],
        }
        count += 1
        time.sleep(2.0)

    if not dry_run and files_meta:
        write_metadata_yaml(folder, "official",
                            "Official kendo organization documents and guidelines",
                            "en", files_meta)
    return count


# ============================================================
# PDF DOWNLOADS
# ============================================================

PDF_SOURCES = [
    {
        "url": "https://industrykendo.com/Articles/Kendo_Training_Handbook.pdf",
        "filename": "kendo-training-handbook.pdf",
        "title": "Kendo Training Handbook by Matt Jackson",
        "tags": ["pdf", "handbook", "training", "instruction"],
    },
    {
        "url": "http://industrykendo.com/Articles/Meaning_of_Kendo_Kata.pdf",
        "filename": "meaning-of-kendo-kata.pdf",
        "title": "Meaning of Kendo Kata",
        "tags": ["pdf", "kata", "philosophy"],
    },
    {
        "url": "http://industrykendo.com/Articles/Fundamental_Theorem_of_Kendo.pdf",
        "filename": "fundamental-theorem-of-kendo.pdf",
        "title": "The Fundamental Theorem of Kendo by Stephen Quinlan",
        "tags": ["pdf", "theory", "philosophy"],
    },
    {
        "url": "https://murdochkendo.com/wp-content/uploads/2023/05/Nihon-Kendo-no-Kata-Kihon-Bokuto-Waza-Kingston-Kendo-Club.pdf",
        "filename": "nihon-kendo-kata-kihon-bokuto-waza.pdf",
        "title": "Nihon Kendo no Kata and Kihon Bokuto Waza Guide",
        "tags": ["pdf", "kata", "instruction"],
    },
    {
        "url": "https://murdochkendo.com/wp-content/uploads/2023/05/The-First-Steps-to-Becoming-a-Referee-Terry-Holt.pdf",
        "filename": "first-steps-becoming-referee.pdf",
        "title": "The First Steps to Becoming a Referee by Terry Holt",
        "tags": ["pdf", "shinpan", "refereeing"],
    },
    {
        "url": "https://murdochkendo.com/wp-content/uploads/2023/05/Kendo-Equipment-Manual-Yasuji-Ishiwata.pdf",
        "filename": "kendo-equipment-manual.pdf",
        "title": "Kendo Equipment Manual by Yasuji Ishiwata",
        "tags": ["pdf", "equipment", "bogu"],
    },
    {
        "url": "http://www.kendo-fik.org/wp-content/uploads/2020/06/STANDARD-GUIDELINE-FOR-DANKYU-EXAMINATION_Kendo_Iaido_Jodo_English.pdf",
        "filename": "fik-dankyu-examination-guideline.pdf",
        "title": "FIK Standard Guideline for Dan/Kyu Examination",
        "tags": ["pdf", "official", "fik", "grading"],
    },
]


def download_pdfs(dry_run: bool = False) -> int:
    print("\n=== Downloading PDF sources ===")
    folder = THEORY_DIR / "pdfs"
    folder.mkdir(parents=True, exist_ok=True)
    files_meta = {}
    count = 0

    for pdf in PDF_SOURCES:
        output_path = folder / pdf["filename"]
        if output_path.exists():
            print(f"  SKIP (exists): {pdf['filename']}")
            continue
        if dry_run:
            print(f"  WOULD DOWNLOAD: {pdf['title']}")
            continue

        print(f"  Downloading: {pdf['title']}...")
        try:
            resp = requests.get(pdf["url"], headers=HEADERS, timeout=60)
            resp.raise_for_status()
            with open(output_path, "wb") as f:
                f.write(resp.content)
            size_kb = len(resp.content) / 1024
            print(f"    Saved: {pdf['filename']} ({size_kb:.0f} KB)")
            files_meta[pdf["filename"]] = {
                "title": pdf["title"],
                "url": pdf["url"],
                "tags": pdf["tags"],
            }
            count += 1
            time.sleep(1.0)
        except Exception as e:
            print(f"    ERROR: {e}")

    if not dry_run and files_meta:
        write_metadata_yaml(folder, "pdfs",
                            "Free kendo PDF handbooks and guides",
                            "en", files_meta)
    return count


# ============================================================
# VIETNAMESE SOURCES
# ============================================================

VIETNAMESE_PAGES = [
    {
        "url": "https://kilala.vn/emagazine/kendo-nghe-thuat-kiem-toi-ren-tinh-than-cua-nguoi-nhat.html",
        "title": "Kendo - Nghe thuat kiem toi ren tinh than cua nguoi Nhat",
        "xpath": '//div[contains(@class, "entry-content")] | //div[contains(@class, "content")] | //article | //div[contains(@class, "article")] | //main',
        "tags": ["vietnamese", "kilala", "history", "philosophy"],
    },
    {
        "url": "https://vjcc.org.vn/cac-linh-vuc-cu-the/kendo-nghe-thuat-kiem-dao-l-u-doi-cua-nhat-ban.html",
        "title": "Kendo - Nghe thuat kiem dao lau doi cua Nhat Ban",
        "xpath": '//div[contains(@class, "entry-content")] | //div[contains(@class, "content")] | //article | //main',
        "tags": ["vietnamese", "vjcc", "history"],
    },
    {
        "url": "https://blog.janbox.com/vi/kendo-la-gi/",
        "title": "Kendo la gi - Lich su phat trien va luat thi dau",
        "xpath": '//div[contains(@class, "entry-content")] | //div[contains(@class, "content")] | //article | //main',
        "tags": ["vietnamese", "janbox", "history", "rules"],
    },
]


def scrape_vietnamese(dry_run: bool = False) -> int:
    print("\n=== Scraping Vietnamese kendo sources ===")
    folder = THEORY_DIR / "blogs" / "vietnamese"
    folder.mkdir(parents=True, exist_ok=True)
    files_meta = {}
    count = 0

    for page in VIETNAMESE_PAGES:
        slug = slugify(page["title"])
        filename = f"{slug}.docx"
        output_path = folder / filename

        if output_path.exists():
            print(f"  SKIP (exists): {filename}")
            continue
        if dry_run:
            print(f"  WOULD SCRAPE: {page['title']}")
            continue

        print(f"  Scraping: {page['title']}...")
        tree = fetch_page(page["url"])
        if tree is None:
            continue

        paragraphs = []
        for xpath in page["xpath"].split(" | "):
            paragraphs = extract_text_elements(tree, xpath)
            if paragraphs:
                break

        if not paragraphs:
            all_p = tree.xpath("//p")
            paragraphs = [p.text_content().strip() for p in all_p
                          if p.text_content().strip() and len(p.text_content().strip()) > 15]

        if not paragraphs:
            print(f"    WARNING: No content for {page['title']}")
            continue

        save_docx(page["title"], paragraphs, output_path, url=page["url"])
        print(f"    Saved: {filename} ({len(paragraphs)} paragraphs)")
        files_meta[filename] = {
            "title": page["title"],
            "url": page["url"],
            "tags": page["tags"],
        }
        count += 1
        time.sleep(2.0)

    if not dry_run and files_meta:
        write_metadata_yaml(folder, "vietnamese",
                            "Vietnamese kendo educational articles",
                            "vi", files_meta)
    return count


# ============================================================
# MAIN
# ============================================================

def main():
    parser = argparse.ArgumentParser(description="Scrape all non-WordPress kendo sources")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--only", choices=["kendoguide", "official", "pdfs", "vietnamese"],
                        help="Only scrape this category")
    args = parser.parse_args()

    total = 0
    if not args.only or args.only == "kendoguide":
        total += scrape_kendoguide(args.dry_run)
    if not args.only or args.only == "official":
        total += scrape_official(args.dry_run)
    if not args.only or args.only == "pdfs":
        total += download_pdfs(args.dry_run)
    if not args.only or args.only == "vietnamese":
        total += scrape_vietnamese(args.dry_run)

    print(f"\n=== Total new items: {total} ===")
    if not args.dry_run and total > 0:
        print("Next steps:")
        print("  1. Review the scraped files")
        print('  2. Run ingestion: ".venv/Scripts/python.exe" scripts/ingest_all.py --reset')


if __name__ == "__main__":
    main()
